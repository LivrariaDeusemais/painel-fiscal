import sys
import os
import json
from unittest.mock import patch
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src" / "codificador"))
from codificador import construir_indices, criar_excel_orcamento, extrair_penkal_orcamento_por_coordenadas, Produto
from openpyxl import load_workbook, Workbook
from codificador import carregar_base, detectar_fornecedor, localizar_anotacoes


class BaseBlingTest(unittest.TestCase):
    def test_atualiza_ean_e_sku_sem_perder_codigo_do_fornecedor(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory) / "base.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Fornecedor", "Código Fornecedor", "EAN", "SKU", "Descrição", "Marca"])
            ws.append(["SBB", "ABC123", "111", "sku1", "Descrição antiga", "SBB"])
            wb.save(base)
            snapshot = Path(directory) / "bling.json"
            snapshot.write_text(json.dumps([{"sku":"SKU1", "ean":"9788531117701", "nome":"Descrição atualizada", "marca":"SBB"}, {"sku":"SKU2", "ean":"222", "nome":"Produto novo", "marca":"CPP"}]))
            with patch.dict(os.environ, {"CODIFICADOR_BLING_SNAPSHOT":str(snapshot)}):
                produtos = carregar_base(base)
            self.assertEqual(len(produtos), 2)
            self.assertEqual(produtos[0].codigo_fornecedor, "ABC123")
            self.assertEqual(produtos[0].fornecedor, "SBB")
            self.assertEqual(produtos[0].ean, "9788531117701")
            self.assertEqual(produtos[0].sku, "SKU1")
            self.assertEqual(produtos[0].descricao, "Descrição atualizada")
            self.assertEqual(construir_indices(produtos)["ean"]["9788531117701"].sku, "SKU1")


    def test_codigo_fornecedor_bling_completa_produto_novo_e_atualiza_existente(self):
        with tempfile.TemporaryDirectory() as directory:
            base = Path(directory) / "base.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Fornecedor", "Código Fornecedor", "EAN", "SKU", "Descrição", "Marca"])
            ws.append(["Penkal", "antigo", "111", "L1349", "Antigo", "Penkal"])
            wb.save(base)
            snapshot = Path(directory) / "bling.json"
            snapshot.write_text(json.dumps([
                {"sku":"L1349", "codigo_fornecedor":"34572", "ean":"9786500000123", "nome":"Período Interbíblico"},
                {"sku":"L9999", "codigo_fornecedor":"00123", "ean":"9786500000456", "nome":"Novo"}]))
            with patch.dict(os.environ, {"CODIFICADOR_BLING_SNAPSHOT":str(snapshot)}):
                indices = construir_indices(carregar_base(base))
            self.assertEqual(indices["codigo"]["34572"].sku, "L1349")
            self.assertEqual(indices["codigo"]["00123"].sku, "L9999")
            self.assertEqual(indices["ean"]["9786500000123"].sku, "L1349")


class CpadTest(unittest.TestCase):
    def test_detecta_razao_social_e_nome_arquivo(self):
        self.assertEqual(detectar_fornecedor("CPAD COMERCIO E DISTRIBUIDORA LTDA"), "cpad")
        self.assertEqual(detectar_fornecedor("Casa Publicadora das Assembleias de Deus"), "cpad")
        self.assertEqual(detectar_fornecedor("", "CPAD-bonificacao.pdf"), "cpad")
        self.assertEqual(detectar_fornecedor("CASA PUBLICADORA PAULISTA"), "cpp")

    def test_nf_multipagina_ignora_cabecalho_e_preserva_faltantes(self):
        def word(code, top):
            return {"text": code, "x0": 7.2, "x1": 52, "top": top, "bottom": top + 7, "height": 7}
        paginas = [{"page": n, "height": 841.89, "words": words} for n, words in (
            (1, [word("1234567890123", 100), word("7898203061842", 491.4)]),
            (2, [word("7898203061842", 272.4), word("7908234013861", 294.7)]),
            (3, [word("7898203061842", 272.4)]))]
        indices = construir_indices([Produto("CPAD", "", "7898203061842", "B1021", "", "")])
        itens, faltantes = localizar_anotacoes(paginas, indices, "cpad", "nf")
        self.assertEqual([i.pagina for i in itens], [1, 2, 3])
        self.assertEqual(faltantes, ["7908234013861"])
        for item in itens:
            self.assertEqual(item.posicao, "abaixo_codigo")
            self.assertEqual(item.x, 7.2)
            top = 491.4 if item.pagina == 1 else 272.4
            baseline = 841.89 - item.y
            self.assertGreater(baseline, top + 7)
            self.assertLess(baseline, top + 22.3)


class OrcamentoPenkalTest(unittest.TestCase):
    def extrair(self, deslocamento=0):
        words = []
        dados = [("34466", "5,0000", "39,9000", "69,9200", "12,0000", "60,00"),
                 ("33616", "5,0000", "64,9000", "64,5600", "23,0000", "115,00"),
                 ("33472", "20,0000", "59,9000", "74,9600", "15,0000", "300,00"),
                 ("31165", "10,0000", "29,9000", "69,9000", "9,0000", "90,00")]
        skus = ["J2006", "L1246", "L1226", "L1236"]
        produtos = [Produto("Penkal", row[0], "", sku, "", "") for row, sku in zip(dados, skus)]
        for i, row in enumerate(dados):
            for text, x in zip((row[0], "UN", *row[1:]), (226, 262, 281, 330, 382, 425, 478)):
                words.append({"text": text, "x0": x + (deslocamento if x > 262 else 0),
                              "top": 274 + i * 24, "bottom": 282 + i * 24})
        return extrair_penkal_orcamento_por_coordenadas(
            [{"words": words, "page": 1, "height": 842}], construir_indices(produtos))

    def test_colunas_liquidas_em_layouts_distintos(self):
        for deslocamento in (0, 40):
            itens = self.extrair(deslocamento)
            self.assertEqual([i.quantidade for i in itens], list(map(Decimal, [5, 5, 20, 10])))
            self.assertEqual([i.valor_unitario for i in itens], list(map(Decimal, [12, 23, 15, 9])))
            self.assertEqual(sum(i.valor_total for i in itens), Decimal(565))
            self.assertEqual([i.sku for i in itens], ["J2006", "L1246", "L1226", "L1236"])

    def test_excel_com_formulas_e_exportacao_padrao_preservada(self):
        with tempfile.TemporaryDirectory() as pasta:
            path = Path(pasta) / "orcamento.xlsx"
            for formulas in (True, False):
                criar_excel_orcamento(path, self.extrair(), formulas=formulas)
                wb = load_workbook(path)
                ws = wb.active
                self.assertEqual([ws.cell(r, 3).value for r in range(2, 6)], [5, 5, 20, 10])
                self.assertEqual(ws["E2"].value, "=D2*C2" if formulas else 60)
                self.assertEqual(ws["E6"].value, "=SUM(E2:E5)" if formulas else 565)
                wb.close()


class OrcamentoCppTest(unittest.TestCase):
    def test_reconhece_modelo_com_logo_em_imagem_e_upload_numerico(self):
        text = "ORÇAMENTO CODIGO DESCRIÇÃO PRODUTO QTDE Vl. Unit Vl. Total Vl. Pedido c/ desconto Nr Pedido de Compra do Cliente"
        self.assertEqual(detectar_fornecedor(text, "1793121771708.pdf"), "cpp")
        self.assertEqual(detectar_fornecedor("ORÇAMENTO CODIGO DESCRIÇÃO QTDE", "upload.pdf"), "auto")

    def pagina(self, number=1, order="1006393", page_label="1 de 1"):
        def word(text,x,top=254,width=20):
            return {"text":text,"x0":x,"x1":x+width,"top":top,"bottom":top+8,"height":8}
        words=[word("7908084623029",30.72,width=58),word("PRETA8",454),word("45,15",510),word("361,20",555)]
        chars=[word("E",465),word("T",470),word("A",475)]
        chars=[{**c,"x1":c["x0"]+5} for c in chars]
        chars += [{**word("8",479),"x1":483}]
        for value,x in [("45,15",510),("361,20",555)]:
            chars += [{**word(ch,x+n*4),"x1":x+n*4+4} for n,ch in enumerate(value)]
        return {"page":number,"height":780,"width":609,"words":words,"chars":chars,
                "text":f"Orçamento Nº {order} Página {page_label}"}

    def test_extrai_colunas_e_anota_na_margem_sem_sobrepor_ean(self):
        from codificador import extrair_cpp_orcamento_por_coordenadas
        from reportlab.pdfbase.pdfmetrics import stringWidth
        indices=construir_indices([Produto("CPP","","7908084623029","B1539","","")])
        pagina=self.pagina()
        item=extrair_cpp_orcamento_por_coordenadas([pagina],indices)[0]
        self.assertEqual((item.quantidade,item.valor_unitario,item.valor_total),(Decimal(8),Decimal('45.15'),Decimal('361.20')))
        anotacoes,_=localizar_anotacoes([pagina],indices,'cpp','orcamento')
        self.assertGreaterEqual(anotacoes[0].x,2)
        self.assertLess(anotacoes[0].x+stringWidth('B1539','Helvetica',6.2),30.72)
        self.assertEqual(anotacoes[0].posicao,'antes_codigo')
        self.assertAlmostEqual(780-anotacoes[0].y,260.2)

    def test_excel_nao_duplica_vias_mas_preserva_pedidos_e_paginas_distintas(self):
        from codificador import extrair_cpp_orcamento_por_coordenadas
        indices=construir_indices([])
        self.assertEqual(len(extrair_cpp_orcamento_por_coordenadas([self.pagina(),self.pagina(2)],indices)),1)
        self.assertEqual(len(extrair_cpp_orcamento_por_coordenadas([self.pagina(),self.pagina(2,order='1006394')],indices)),2)
        self.assertEqual(len(extrair_cpp_orcamento_por_coordenadas([self.pagina(page_label='1 de 2'),self.pagina(2,page_label='2 de 2')],indices)),2)


if __name__ == "__main__":
    unittest.main()
