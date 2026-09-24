from __future__ import annotations

import io
import json
import os
import re
import shutil
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

import pdfplumber
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader, PdfWriter
from reportlab.lib.colors import red
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parent
PACKAGED_DATA_DIR = ROOT / "data"
STORAGE_ROOT = Path(os.environ.get("CODIFICADOR_STORAGE_DIR", ROOT / "runtime")).resolve()
DATA_DIR = STORAGE_ROOT / "data"
UPLOADS_DIR = STORAGE_ROOT / "uploads"
OUTPUTS_DIR = STORAGE_ROOT / "outputs"
DEFAULT_BASE = PACKAGED_DATA_DIR / "base_skus.xlsx"
BASE_FILE = DATA_DIR / "base_skus.xlsx"
MODELO_BASE_FILE = PACKAGED_DATA_DIR / "modelo_base_skus.xlsx"

for folder in (DATA_DIR, UPLOADS_DIR, OUTPUTS_DIR):
    folder.mkdir(parents=True, exist_ok=True)


@dataclass
class Produto:
    fornecedor: str
    codigo_fornecedor: str
    ean: str
    sku: str
    descricao: str
    marca: str


@dataclass
class ItemPdf:
    codigo: str
    sku: str
    quantidade: Decimal | None
    valor_total: Decimal | None
    valor_unitario: Decimal | None
    pagina: int
    x: float
    y: float
    texto: str
    posicao: str


def normalizar_texto(valor: object) -> str:
    texto = "" if valor is None else str(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = texto.lower().strip()
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    return re.sub(r"\s+", " ", texto).strip()


def limpar(valor: object) -> str:
    if valor is None:
        return ""
    texto = str(valor).strip().strip("\t\r\n ")
    return "" if texto.lower() == "nan" else texto


def apenas_digitos(valor: object) -> str:
    return re.sub(r"\D+", "", limpar(valor))


def decimal_br(valor: str) -> Decimal | None:
    texto = limpar(valor).replace("R$", "").replace(" ", "")
    if not texto:
        return None
    texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto)
    except Exception:
        return None


def decimal_excel(valor: Decimal | None) -> float | None:
    if valor is None:
        return None
    return float(valor.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def arquivo_base_ativo() -> Path:
    return BASE_FILE if BASE_FILE.exists() else DEFAULT_BASE


def atualizar_base(novo_arquivo: Path) -> Path:
    carregar_base(novo_arquivo)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    arquivo_temporario = BASE_FILE.with_suffix(".tmp.xlsx")
    shutil.copy2(novo_arquivo, arquivo_temporario)
    arquivo_temporario.replace(BASE_FILE)
    return BASE_FILE


def criar_modelo_base() -> Path:
    if MODELO_BASE_FILE.exists():
        return MODELO_BASE_FILE
    PACKAGED_DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Base de SKUs"
    headers = ["Fornecedor", "Cód. no fornecedor", "GTIN/EAN", "SKU Deus é Mais", "Descrição", "Marca"]
    ws.append(headers)
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.alignment = Alignment(horizontal="center")
    widths = [22, 24, 20, 20, 52, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"

    instrucoes = wb.create_sheet("Instruções")
    instrucoes["A1"] = "Como atualizar a base de dados"
    instrucoes["A1"].font = Font(bold=True, size=14)
    instrucoes["A3"] = (
        'Para obter a base de dados em Excel, acessar o Bling em CADASTRO/PRODUTO, clicar em '
        '"exportar dados para planilha" e utilizar as colunas correspondente do arquivo do Bling '
        "para preencher a planilha modelo."
    )
    instrucoes["A5"] = "Colunas obrigatórias"
    instrucoes["A5"].font = Font(bold=True)
    for linha, texto in enumerate(
        [
            "Fornecedor: nome da editora/fornecedor, como SBB, BV Books, CPP ou Penkal.",
            "Cód. no fornecedor: código interno usado pelo fornecedor, quando existir.",
            "GTIN/EAN: EAN ou ISBN do produto, somente números quando possível.",
            "SKU Deus é Mais: SKU cadastrado na Deus é Mais.",
            "Descrição: nome/descrição do produto.",
            "Marca: marca ou editora do produto.",
        ],
        start=6,
    ):
        instrucoes.cell(linha, 1, texto)
    instrucoes.column_dimensions["A"].width = 115
    wb.save(MODELO_BASE_FILE)
    return MODELO_BASE_FILE


def carregar_base(path: Path | None = None) -> list[Produto]:
    path = path or arquivo_base_ativo()
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        raise ValueError("A base de SKUs está vazia.")

    headers = [normalizar_texto(celula) for celula in linhas[0]]

    def achar(*nomes: str) -> int:
        nomes_norm = {normalizar_texto(nome) for nome in nomes}
        for indice, header in enumerate(headers):
            if header in nomes_norm:
                return indice
        for indice, header in enumerate(headers):
            if any(nome in header for nome in nomes_norm):
                return indice
        raise ValueError(f"Coluna não encontrada na base: {nomes}")

    col_fornecedor = achar("Fornecedor")
    col_codigo = achar("Cód. no fornecedor", "Código Fornecedor")
    col_ean = achar("GTIN/EAN", "EAN", "ISBN")
    col_sku = achar("SKU Deus é Mais", "SKU")
    col_descricao = achar("Descrição")
    col_marca = achar("Marca")

    produtos: list[Produto] = []
    for linha in linhas[1:]:
        sku = limpar(linha[col_sku] if col_sku < len(linha) else "")
        if not sku:
            continue
        produtos.append(
            Produto(
                fornecedor=limpar(linha[col_fornecedor] if col_fornecedor < len(linha) else ""),
                codigo_fornecedor=limpar(linha[col_codigo] if col_codigo < len(linha) else ""),
                ean=apenas_digitos(linha[col_ean] if col_ean < len(linha) else ""),
                sku=sku,
                descricao=limpar(linha[col_descricao] if col_descricao < len(linha) else ""),
                marca=limpar(linha[col_marca] if col_marca < len(linha) else ""),
            )
        )
    return produtos


def construir_indices(produtos: Iterable[Produto]) -> dict[str, dict[str, Produto]]:
    por_ean: dict[str, Produto] = {}
    por_codigo: dict[str, Produto] = {}
    for produto in produtos:
        if produto.ean:
            por_ean.setdefault(produto.ean, produto)
        if produto.codigo_fornecedor:
            por_codigo.setdefault(produto.codigo_fornecedor, produto)
            dig = apenas_digitos(produto.codigo_fornecedor)
            if dig:
                por_codigo.setdefault(dig, produto)
    return {"ean": por_ean, "codigo": por_codigo}


def texto_pdf(pdf_path: Path) -> str:
    partes: list[str] = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            partes.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
    return "\n".join(partes)


def detectar_fornecedor(texto: str, nome_arquivo: str = "") -> str:
    base = normalizar_texto(texto)
    nome = normalizar_texto(nome_arquivo)
    if "sociedade biblica do brasil" in base:
        return "sbb"
    if "bkj1611" in base or "biblia king james" in base:
        return "bv_books"
    if "casa publicadora paulista" in base:
        return "cpp"
    if "editora penkal" in base:
        return "penkal"
    if "adib editora" in base or "inteligencia biblica" in base:
        return "adib"
    if "sbb" in nome:
        return "sbb"
    if "bv books" in nome or "bkj1611" in nome:
        return "bv_books"
    if "cpp" in nome:
        return "cpp"
    if "penkal" in nome:
        return "penkal"
    if "adib" in nome or "inteligencia biblica" in nome:
        return "adib"
    return "auto"


def detectar_tipo(texto: str) -> str:
    base = normalizar_texto(texto)
    if "orcamento" in base or "proposta comercial" in base or "pedido de venda" in base:
        return "orcamento"
    return "nf"


def extrair_paginas(pdf_path: Path):
    paginas = []
    with pdfplumber.open(pdf_path) as pdf:
        for index, page in enumerate(pdf.pages, 1):
            paginas.append(
                {
                    "page": index,
                    "width": page.width,
                    "height": page.height,
                    "rotation": page.rotation,
                    "words": page.extract_words(
                        x_tolerance=1.2,
                        y_tolerance=3,
                        keep_blank_chars=False,
                        use_text_flow=False,
                    ),
                    "text": page.extract_text(x_tolerance=1.5, y_tolerance=3) or "",
                }
            )
    return paginas


def agrupar_linhas(words: list[dict]) -> list[list[dict]]:
    linhas: list[list[dict]] = []
    for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
        centro = (word["top"] + word["bottom"]) / 2
        for linha in linhas:
            centro_linha = sum((w["top"] + w["bottom"]) / 2 for w in linha) / len(linha)
            if abs(centro - centro_linha) <= 3:
                linha.append(word)
                break
        else:
            linhas.append([word])
    for linha in linhas:
        linha.sort(key=lambda item: item["x0"])
    return linhas


def texto_bloco(words: list[dict]) -> str:
    return " ".join(w["text"] for w in sorted(words, key=lambda item: (item["top"], item["x0"])))


def buscar_produto(codigo: str, bloco: str, indices: dict[str, dict[str, Produto]]) -> Produto | None:
    codigo_limpo = limpar(codigo)
    dig = apenas_digitos(codigo_limpo)
    for chave in (codigo_limpo, dig):
        if chave and chave in indices["codigo"]:
            return indices["codigo"][chave]
        if chave and chave in indices["ean"]:
            return indices["ean"][chave]
    isbn = re.search(r"ISBN[:\s]*([0-9]{10,13})", bloco, re.IGNORECASE)
    if isbn and isbn.group(1) in indices["ean"]:
        return indices["ean"][isbn.group(1)]
    return None


def extrair_itens_por_linha(texto: str, indices: dict[str, dict[str, Produto]], fornecedor: str, tipo: str) -> list[ItemPdf]:
    itens: list[ItemPdf] = []
    linhas = [linha.strip() for linha in texto.splitlines() if linha.strip()]
    i = 0
    while i < len(linhas):
        linha = linhas[i]
        codigo = ""
        if fornecedor == "bv_books":
            m = re.match(r"^((?:BL|LV)-\d+)\s+(.+)", linha)
        elif fornecedor == "penkal":
            m = re.match(r"^(.+?)\s+(\d{4,8})\s+UN\s+([\d,.]+)\s+", linha)
            if m:
                codigo = m.group(2)
        else:
            m = re.match(r"^(\d{10,13})\s+(.+)", linha)
        if not m:
            i += 1
            continue
        if not codigo:
            codigo = m.group(1)
        bloco = linha
        j = i + 1
        while j < len(linhas) and not re.match(r"^((?:BL|LV)-\d+|\d{10,13}|.+?\s+\d{4,8}\s+UN\s+[\d,.]+\s+)", linhas[j]):
            if any(final in normalizar_texto(linhas[j]) for final in ("subtotal", "qt total", "total dos produtos")):
                break
            bloco += " " + linhas[j]
            j += 1
        produto = buscar_produto(codigo, bloco, indices)
        quantidade, total = extrair_quantidade_total(bloco, fornecedor, tipo)
        unitario = None
        if quantidade and total and quantidade != 0:
            unitario = (total / quantidade).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        itens.append(
            ItemPdf(
                codigo=codigo,
                sku=produto.sku if produto else "",
                quantidade=quantidade,
                valor_total=total,
                valor_unitario=unitario,
                pagina=1,
                x=0,
                y=0,
                texto=bloco,
                posicao="excel",
            )
        )
        i = max(j, i + 1)
    return itens


def extrair_quantidade_total(bloco: str, fornecedor: str, tipo: str) -> tuple[Decimal | None, Decimal | None]:
    numeros = re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2,4}|\d+", bloco)
    if not numeros:
        return None, None
    if fornecedor == "bv_books" and tipo == "orcamento":
        candidatos = re.search(r"\b(\d{1,4},00)\s+UN\s+[\d.,]+\s+[\d.,]+\s+([\d.]+,\d{2})", bloco)
        if candidatos:
            return decimal_br(candidatos.group(1)), decimal_br(candidatos.group(2))
    if fornecedor == "sbb" and tipo == "orcamento":
        candidatos = re.search(r"\s(\d+)\s+R\$\s+[\d.,]+\s+R\$\s+[\d.,]+\s+\d+%\s+R\$\s+([\d.,]+)", bloco)
        if candidatos:
            return decimal_br(candidatos.group(1)), decimal_br(candidatos.group(2))
    if fornecedor == "cpp" and tipo == "orcamento":
        candidatos = re.search(r"\s(\d+)\s+([\d.,]+)\s+([\d.,]+)$", bloco)
        if candidatos:
            return decimal_br(candidatos.group(1)), decimal_br(candidatos.group(3))
    if fornecedor == "penkal" and tipo == "orcamento":
        candidatos = re.search(r"\sUN\s+([\d.,]+)\s+[\d.,]+\s+[\d.,]+\s+[\d.,]+\s+([\d.,]+)$", bloco)
        if candidatos:
            return decimal_br(candidatos.group(1)), decimal_br(candidatos.group(2))
    return None, None


def localizar_anotacoes(paginas, indices, fornecedor: str, tipo: str) -> tuple[list[ItemPdf], list[str]]:
    if fornecedor == "bv_books" and tipo == "orcamento":
        return localizar_bv_orcamento(paginas, indices)
    if fornecedor == "bv_books":
        return localizar_bv(paginas, indices, tipo)
    if tipo == "orcamento" and fornecedor == "sbb":
        return localizar_generico_ean(
            paginas,
            indices,
            font_size=8,
            min_x=25,
            max_x=145,
            min_top=180,
            continuation_min_top=70,
            max_height=9,
            antes_codigo=True,
        )
    if tipo == "orcamento" and fornecedor == "cpp":
        return localizar_generico_ean(
            paginas,
            indices,
            font_size=6.2,
            min_x=25,
            max_x=145,
            min_top=180,
            max_height=9,
            antes_codigo_gap=8,
            antes_codigo=True,
        )
    if tipo == "orcamento" and fornecedor == "penkal":
        return localizar_penkal_orcamento(paginas, indices)
    if fornecedor == "sbb":
        return localizar_generico_ean(paginas, indices, font_size=9, min_x=75, max_x=130, min_top=140, max_height=8, apos_descricao=True)
    if fornecedor == "cpp":
        return localizar_generico_ean(
            paginas,
            indices,
            font_size=6.2,
            min_x=0,
            max_x=60,
            min_top=450,
            max_height=8,
            continuation_min_top=260,
            abaixo_codigo=True,
        )
    if fornecedor == "penkal":
        return localizar_penkal(paginas, indices)
    if fornecedor == "adib":
        return localizar_adib(paginas, indices)
    return localizar_generico_ean(paginas, indices, font_size=8, min_x=0, max_x=150, min_top=180, max_height=9, abaixo_codigo=True)


def localizar_bv(paginas, indices, tipo: str) -> tuple[list[ItemPdf], list[str]]:
    anotacoes: list[ItemPdf] = []
    faltantes: list[str] = []
    code_re = re.compile(r"^(BL|LV)-\d+$")
    for pagina in paginas:
        words = pagina["words"]
        code_words = sorted([w for w in words if code_re.match(w["text"])], key=lambda item: (item["top"], item["x0"]))
        for idx, code_word in enumerate(code_words):
            next_top = code_words[idx + 1]["top"] if idx + 1 < len(code_words) else code_word["top"] + 40
            bloco_words = [w for w in words if code_word["top"] - 1 <= w["top"] < next_top - 1 and w["x0"] > code_word["x1"]]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            if not produto:
                faltantes.append(code_word["text"])
                continue
            linhas_baixo = [linha for linha in agrupar_linhas(bloco_words) if linha[0]["top"] > code_word["top"] + 2]
            if linhas_baixo:
                x = code_word["x0"] + 0.5
                y = pagina["height"] - linhas_baixo[0][0]["bottom"] - 0.2
                posicao = "abaixo_codigo"
            else:
                isbn_words = [w for w in bloco_words if "ISBN" in w["text"].upper()]
                alvo = max(isbn_words or bloco_words, key=lambda item: (item["top"], item["x1"]))
                ncm_words = [w for w in bloco_words if re.fullmatch(r"\d{4}\.\d{2}\.\d{2}|\d{8}", w["text"])]
                limite = min((w["x0"] for w in ncm_words), default=pagina["width"] - 30)
                largura = stringWidth(produto.sku, "Helvetica", 8)
                x = min(alvo["x1"] + 3, limite - largura - 2)
                y = pagina["height"] - alvo["bottom"] - 0.2
                posicao = "apos_isbn"
            anotacoes.append(ItemPdf(code_word["text"], produto.sku, None, None, None, pagina["page"], x, y, bloco, posicao))
    return anotacoes, faltantes


def localizar_bv_orcamento(paginas, indices) -> tuple[list[ItemPdf], list[str]]:
    anotacoes: list[ItemPdf] = []
    faltantes: list[str] = []
    code_re = re.compile(r"^(BL|LV)-\d+$")
    for pagina in paginas:
        words = pagina["words"]
        code_words = sorted([w for w in words if code_re.match(w["text"])], key=lambda item: (item["top"], item["x0"]))
        for idx, code_word in enumerate(code_words):
            next_top = code_words[idx + 1]["top"] if idx + 1 < len(code_words) else code_word["top"] + 40
            bloco_words = [w for w in words if code_word["top"] - 1 <= w["top"] < next_top - 1 and w["x0"] > code_word["x1"]]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            if not produto:
                faltantes.append(code_word["text"])
                continue
            anotacoes.append(
                ItemPdf(
                    code_word["text"],
                    produto.sku,
                    None,
                    None,
                    None,
                    pagina["page"],
                    code_word["x1"] + 5,
                    pagina["height"] - code_word["bottom"] + 1.5,
                    bloco,
                    "entre_codigo_descricao",
                )
            )
    return anotacoes, faltantes


def localizar_generico_ean(
    paginas,
    indices,
    font_size: float,
    min_x: float,
    max_x: float,
    min_top: float,
    max_height: float,
    continuation_min_top: float | None = None,
    apos_descricao: bool = False,
    abaixo_codigo: bool = False,
    apos_codigo: bool = False,
    antes_codigo: bool = False,
    antes_codigo_gap: float = 3,
) -> tuple[list[ItemPdf], list[str]]:
    anotacoes: list[ItemPdf] = []
    faltantes: list[str] = []
    for pagina in paginas:
        words = pagina["words"]
        codigos = sorted(
            [
                w
                for w in words
                if re.fullmatch(r"\d{10,13}", w["text"])
                and min_x <= w["x0"] <= max_x
                and w["top"] >= (continuation_min_top if pagina["page"] > 1 and continuation_min_top is not None else min_top)
                and w["height"] <= max_height
            ],
            key=lambda item: (item["top"], item["x0"]),
        )
        for idx, code_word in enumerate(codigos):
            next_top = codigos[idx + 1]["top"] if idx + 1 < len(codigos) else code_word["top"] + 35
            bloco_words = [w for w in words if code_word["top"] - 1 <= w["top"] < next_top - 1 and w["x0"] > code_word["x1"]]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            if not produto:
                faltantes.append(code_word["text"])
                continue
            if antes_codigo:
                largura = stringWidth(produto.sku, "Helvetica", font_size)
                x = max(2, code_word["x0"] - largura - antes_codigo_gap)
                y = pagina["height"] - code_word["bottom"] + 1.8
                posicao = "antes_codigo"
            elif apos_descricao:
                ncm = next((w for w in bloco_words if re.fullmatch(r"\d{8}", w["text"])), None)
                linha = [w for w in bloco_words if abs(w["top"] - code_word["top"]) <= 3]
                desc = [w for w in linha if not ncm or w["x1"] < ncm["x0"]]
                alvo = max(desc or linha or [code_word], key=lambda w: w["x1"])
                limite = ncm["x0"] if ncm else pagina["width"] - 25
                largura = stringWidth(produto.sku, "Helvetica", font_size)
                x = min(alvo["x1"] + 3, limite - largura - 2)
                y = pagina["height"] - alvo["bottom"] - 0.2
                posicao = "apos_descricao"
            elif abaixo_codigo:
                x = code_word["x0"]
                y = pagina["height"] - code_word["bottom"] - font_size * 0.95
                posicao = "abaixo_codigo"
            elif apos_codigo:
                x = code_word["x1"] + 3
                y = pagina["height"] - code_word["bottom"] - 0.2
                posicao = "apos_codigo"
            else:
                x = code_word["x1"] + 3
                y = pagina["height"] - code_word["bottom"] - 0.2
                posicao = "apos_codigo"
            anotacoes.append(ItemPdf(code_word["text"], produto.sku, None, None, None, pagina["page"], x, y, bloco, posicao))
    return anotacoes, faltantes


def localizar_penkal(paginas, indices) -> tuple[list[ItemPdf], list[str]]:
    anotacoes: list[ItemPdf] = []
    faltantes: list[str] = []
    for pagina in paginas:
        words = pagina["words"]
        min_top = 180 if pagina["page"] > 1 else 250
        codigos = sorted(
            [w for w in words if re.fullmatch(r"\d{4,8}", w["text"]) and 20 <= w["x0"] <= 90 and w["top"] > min_top],
            key=lambda item: (item["top"], item["x0"]),
        )
        for code_word in codigos:
            bloco_words = [w for w in words if abs(w["top"] - code_word["top"]) <= 14 and w["x0"] > code_word["x1"]]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            if not produto:
                faltantes.append(code_word["text"])
                continue
            anotacoes.append(
                ItemPdf(
                    code_word["text"],
                    produto.sku,
                    None,
                    None,
                    None,
                    pagina["page"],
                    9.8,
                    pagina["height"] - code_word["bottom"] + 1.85,
                    bloco,
                    "margem_esquerda",
                )
            )
    return anotacoes, faltantes


def localizar_adib(paginas, indices) -> tuple[list[ItemPdf], list[str]]:
    anotacoes: list[ItemPdf] = []
    faltantes: list[str] = []
    for pagina in paginas:
        words = pagina["words"]
        codigos = sorted(
            [
                w
                for w in words
                if re.fullmatch(r"\d{1,5}", w["text"])
                and 70 <= w["x0"] <= 85
                and w["top"] > 330
                and w["height"] <= 9
            ],
            key=lambda item: (item["top"], item["x0"]),
        )
        for idx, code_word in enumerate(codigos):
            next_top = codigos[idx + 1]["top"] if idx + 1 < len(codigos) else code_word["top"] + 35
            bloco_words = [w for w in words if code_word["top"] - 1 <= w["top"] < next_top - 1 and w["x0"] > code_word["x1"]]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            if not produto:
                faltantes.append(code_word["text"])
                continue
            largura = stringWidth(produto.sku, "Helvetica", 8)
            anotacoes.append(
                ItemPdf(
                    code_word["text"],
                    produto.sku,
                    None,
                    None,
                    None,
                    pagina["page"],
                    max(38, code_word["x0"] - largura - 5),
                    pagina["height"] - code_word["bottom"] + 1.5,
                    bloco,
                    "codigo_adib",
                )
            )
    return anotacoes, faltantes


def localizar_penkal_orcamento(paginas, indices) -> tuple[list[ItemPdf], list[str]]:
    anotacoes: list[ItemPdf] = []
    faltantes: list[str] = []
    for item in extrair_penkal_orcamento_por_coordenadas(paginas, indices):
        if not item.sku:
            faltantes.append(item.codigo)
            continue
        item.x = 10.5
        item.y = item.y + 2.05
        item.posicao = "margem_esquerda_descricao"
        anotacoes.append(item)
    return anotacoes, faltantes


def extrair_penkal_orcamento_por_coordenadas(paginas, indices) -> list[ItemPdf]:
    itens: list[ItemPdf] = []
    for pagina in paginas:
        words = pagina["words"]
        codigos = sorted(
            [
                w
                for w in words
                if re.fullmatch(r"\d{4,8}", w["text"])
                and 180 <= w["x0"] <= 285
                and (w["top"] > 230 or pagina["page"] > 1)
            ],
            key=lambda item: (item["top"], item["x0"]),
        )
        for idx, code_word in enumerate(codigos):
            next_top = codigos[idx + 1]["top"] if idx + 1 < len(codigos) else code_word["top"] + 35
            bloco_words = [w for w in words if code_word["top"] - 24 <= w["top"] < next_top - 1]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            same_line = [w for w in words if abs(w["top"] - code_word["top"]) <= 2.5]
            quantidade_word = next((w for w in same_line if 300 <= w["x0"] <= 350), None)
            total_word = next((w for w in same_line if 515 <= w["x0"] <= 560), None)
            quantidade = decimal_br(quantidade_word["text"]) if quantidade_word else None
            total = decimal_br(total_word["text"]) if total_word else None
            unitario = None
            if quantidade and total and quantidade != 0:
                unitario = (total / quantidade).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            itens.append(
                ItemPdf(
                    code_word["text"],
                    produto.sku if produto else "",
                    quantidade,
                    total,
                    unitario,
                    pagina["page"],
                    10.5,
                    pagina["height"] - code_word["bottom"] - 0.2,
                    bloco,
                    "excel",
                )
            )
    return itens


def extrair_decimal_de_texto(valor: str) -> Decimal | None:
    candidatos = re.findall(r"\d{1,3}(?:\.\d{3})*,\d{2}", limpar(valor))
    if not candidatos:
        return None
    return decimal_br(candidatos[-1])


def extrair_cpp_orcamento_por_coordenadas(paginas, indices) -> list[ItemPdf]:
    itens: list[ItemPdf] = []
    for pagina in paginas:
        words = pagina["words"]
        codigos = sorted(
            [
                w
                for w in words
                if re.fullmatch(r"\d{10,13}", w["text"])
                and 25 <= w["x0"] <= 45
                and w["top"] >= 240
                and w["height"] <= 9
            ],
            key=lambda item: (item["top"], item["x0"]),
        )
        for idx, code_word in enumerate(codigos):
            next_top = codigos[idx + 1]["top"] if idx + 1 < len(codigos) else code_word["top"] + 25
            bloco_words = [w for w in words if code_word["top"] - 1 <= w["top"] < next_top - 1]
            bloco = texto_bloco(bloco_words)
            produto = buscar_produto(code_word["text"], bloco, indices)
            same_line = [w for w in words if abs(w["top"] - code_word["top"]) <= 2.5]

            quantidade = None
            for word in sorted(same_line, key=lambda item: item["x0"]):
                if 455 <= word["x0"] <= 505:
                    if re.fullmatch(r"\d+", word["text"]):
                        quantidade = decimal_br(word["text"])
                        break
                    digitos = re.sub(r"\D+", "", word["text"])
                    if "," not in word["text"] and len(digitos) == 1:
                        quantidade = decimal_br(digitos)
                        break

            total_word = next((w for w in same_line if 545 <= w["x0"] <= 585 and "," in w["text"]), None)
            total = extrair_decimal_de_texto(total_word["text"]) if total_word else None
            unitario = None
            if quantidade and total and quantidade != 0:
                unitario = (total / quantidade).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            itens.append(
                ItemPdf(
                    code_word["text"],
                    produto.sku if produto else "",
                    quantidade,
                    total,
                    unitario,
                    pagina["page"],
                    0,
                    0,
                    bloco,
                    "excel",
                )
            )
    return itens


def overlay_pdf(input_pdf: Path, output_pdf: Path, paginas, anotacoes: list[ItemPdf], font_size: float) -> None:
    por_pagina: dict[int, list[ItemPdf]] = defaultdict(list)
    for item in anotacoes:
        por_pagina[item.pagina].append(item)

    reader = PdfReader(str(input_pdf))
    writer = PdfWriter()
    for idx, page in enumerate(reader.pages, 1):
        page.transfer_rotation_to_content()
        itens = por_pagina.get(idx, [])
        if itens:
            info = paginas[idx - 1]
            packet = io.BytesIO()
            c = canvas.Canvas(packet, pagesize=(info["width"], info["height"]))
            c.setFillColor(red)
            c.setFont("Helvetica", font_size)
            for item in itens:
                c.drawString(item.x, item.y, item.sku)
            c.save()
            packet.seek(0)
            page.merge_page(PdfReader(packet).pages[0])
        writer.add_page(page)
    with output_pdf.open("wb") as handle:
        writer.write(handle)


def criar_excel_orcamento(output_xlsx: Path, itens: list[ItemPdf]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orçamento"
    headers = ["Código do fornecedor", "SKU Deus é Mais", "Quantidade orçada", "Valor unitário", "Valor total"]
    ws.append(headers)
    total_geral = Decimal("0")
    for item in itens:
        total_item = item.valor_total
        if total_item is None and item.quantidade is not None and item.valor_unitario is not None:
            total_item = (item.quantidade * item.valor_unitario).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if total_item is not None:
            total_geral += total_item
        ws.append(
            [
                item.codigo,
                item.sku,
                decimal_excel(item.quantidade),
                decimal_excel(item.valor_unitario),
                decimal_excel(total_item),
            ]
        )
    total_row = ws.max_row + 1
    ws.cell(total_row, 4, "Total geral")
    ws.cell(total_row, 5, decimal_excel(total_geral))
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")
        celula.alignment = Alignment(horizontal="center")
    for celula in ws[total_row]:
        celula.font = Font(bold=True)
    widths = [22, 18, 20, 18, 18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = '#,##0.00'
    wb.save(output_xlsx)


def processar(pdf_path: Path, fornecedor: str = "auto", tipo: str = "auto") -> dict:
    texto = texto_pdf(pdf_path)
    fornecedor_final = detectar_fornecedor(texto, pdf_path.name) if fornecedor == "auto" else fornecedor
    tipo_final = detectar_tipo(texto) if tipo == "auto" else tipo
    produtos = carregar_base()
    indices = construir_indices(produtos)
    paginas = extrair_paginas(pdf_path)
    anotacoes, faltantes = localizar_anotacoes(paginas, indices, fornecedor_final, tipo_final)
    itens_excel: list[ItemPdf] = []
    if tipo_final == "orcamento":
        if fornecedor_final == "penkal":
            itens_excel = extrair_penkal_orcamento_por_coordenadas(paginas, indices)
        elif fornecedor_final == "cpp":
            itens_excel = extrair_cpp_orcamento_por_coordenadas(paginas, indices)
        else:
            itens_excel = extrair_itens_por_linha(texto, indices, fornecedor_final, tipo_final)
        if itens_excel:
            by_code = {item.codigo: item for item in itens_excel}
            for item in anotacoes:
                if item.codigo in by_code:
                    item.quantidade = by_code[item.codigo].quantidade
                    item.valor_total = by_code[item.codigo].valor_total
                    item.valor_unitario = by_code[item.codigo].valor_unitario

    limite_saida = datetime.now().timestamp() - (7 * 24 * 60 * 60)
    for pasta_antiga in OUTPUTS_DIR.iterdir():
        if pasta_antiga.is_dir() and pasta_antiga.stat().st_mtime < limite_saida:
            shutil.rmtree(pasta_antiga, ignore_errors=True)

    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    pasta_saida = OUTPUTS_DIR / timestamp
    pasta_saida.mkdir(parents=True, exist_ok=True)
    base_nome = re.sub(r"[^A-Za-z0-9_-]+", "_", pdf_path.stem).strip("_")
    pdf_saida = pasta_saida / f"{base_nome}_com_sku.pdf"
    font_size = 9 if fornecedor_final == "sbb" and tipo_final == "nf" else 8
    if fornecedor_final == "cpp":
        font_size = 6.2
    if fornecedor_final == "penkal":
        font_size = 6.5
    overlay_pdf(pdf_path, pdf_saida, paginas, anotacoes, font_size)

    excel_saida = None
    if tipo_final == "orcamento":
        excel_saida = pasta_saida / f"{base_nome}_itens.xlsx"
        criar_excel_orcamento(excel_saida, itens_excel or anotacoes)

    total_itens = len(itens_excel) if tipo_final == "orcamento" and itens_excel else len(anotacoes)
    resumo = {
        "fornecedor": fornecedor_final,
        "tipo": tipo_final,
        "pdf": str(pdf_saida),
        "excel": str(excel_saida) if excel_saida else None,
        "itens_encontrados": total_itens,
        "skus_aplicados": len(anotacoes),
        "sem_correspondencia": sorted(set(faltantes)),
        "base": str(arquivo_base_ativo()),
    }
    (pasta_saida / "resumo.json").write_text(json.dumps(resumo, ensure_ascii=False, indent=2), encoding="utf-8")
    return resumo
