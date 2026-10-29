"""Read-only imports for inventory planning. Never executes workbook macros/formulas."""
import sys, json, re
from pathlib import Path
from decimal import Decimal
import pdfplumber
import openpyxl
import xlrd

def number(v):
    if isinstance(v, (int, float)): return float(v)
    s = str(v or '').strip().replace('R$', '').replace(' ', '')
    if ',' in s: s = s.replace('.', '').replace(',', '.')
    try: return float(Decimal(s))
    except Exception: raise ValueError('Quantidade ou valor inválido: ' + str(v))

def norm(v):
    import unicodedata
    return ''.join(c for c in unicodedata.normalize('NFD', str(v or '').lower().strip()) if unicodedata.category(c) != 'Mn')

def spreadsheet(file, kind):
    if Path(file).suffix.lower() == '.xls':
        book = xlrd.open_workbook(file, ignore_workbook_corruption=True, logfile=sys.stderr)
        sheet = book.sheet_by_index(0)
        if sheet.nrows>20000 or sheet.ncols>100: raise ValueError('Planilha excede 20 mil linhas ou 100 colunas.')
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
    else:
        book = openpyxl.load_workbook(file, data_only=True, read_only=True)
        sheet=book.worksheets[0]
        if sheet.max_row>20000 or sheet.max_column>100: raise ValueError('Planilha excede 20 mil linhas ou 100 colunas.')
        rows = list(sheet.values)
        book.close()
    if len(rows) > 20000: raise ValueError('Limite de 20 mil linhas por arquivo.')
    required = ['codigo', 'quantidade', 'cfop'] if kind == 'giro' else ['sku deus e mais', 'quantidade orcada', 'valor unitario']
    start = next((i for i,r in enumerate(rows[:20]) if all(h in [norm(c) for c in r] for h in required)), None)
    if start is None: raise ValueError('Cabeçalhos não reconhecidos. Use o relatório por produto do Bling ou o Excel do Codificador.')
    headers = [norm(c) for c in rows[start]]
    def get(row, key): return row[headers.index(key)] if key in headers and headers.index(key) < len(row) else None
    out = []
    for i, row in enumerate(rows[start+1:], start+2):
        sku = str(get(row, 'codigo' if kind == 'giro' else 'sku deus e mais') or '').strip()
        code = str(get(row, 'codigo do fornecedor') or '').strip()
        if not sku and not code: continue
        if kind == 'giro':
            cfop = str(int(number(get(row, 'cfop'))))
            if not re.fullmatch(r'[567]\d{3}', cfop): raise ValueError('CFOP de saída inválido na linha '+str(i))
            out.append({'sku':sku,'quantity':number(get(row,'quantidade')),'cfop':cfop,'line':i})
        else:
            out.append({'sku':sku,'code':code,'quantity':number(get(row,'quantidade orcada')),'price':number(get(row,'valor unitario')),'line':i})
    if not out: raise ValueError('Nenhum item encontrado.')
    return {'rows':out}

def full(file, channel):
    out, shipment, declared = [], '', None
    with pdfplumber.open(file) as pdf:
        if len(pdf.pages)>100: raise ValueError('PDF com mais de 100 páginas.')
        for page in pdf.pages:
            text = page.extract_text() or ''
            if channel == 'meli':
                m = re.search(r'Frete\s*#\s*(\d+)',text)
                if m: shipment=m[1]
                m = re.search(r'Total de unidades:\s*(\d+)',text)
                if m: declared=int(m[1])
                words=page.extract_words()
                anchors=[w for w in words if w['text']=='Código' and w['x0']<page.width*.12 and any(x['text'].startswith('ML:') and abs(x['top']-w['top'])<3 for x in words)]
                for i,a in enumerate(anchors):
                    bottom=anchors[i+1]['top']-5 if i+1<len(anchors) else page.height
                    block=[w for w in words if a['top']-5<=w['top']<bottom]
                    joined=' '.join(w['text'] for w in block)
                    sku=re.search(r'SKU:\s*([A-Za-z0-9_-]+)',joined)
                    qty=[w['text'] for w in block if page.width*.37<w['x0']<page.width*.47 and re.fullmatch(r'\d+',w['text'])]
                    ean=re.search(r'(\d{12,14})',joined)
                    if not sku or len(qty)!=1: raise ValueError('Não foi possível separar SKU e quantidade no PDF Mercado Livre.')
                    out.append({'sku':sku[1],'ean':ean[1] if ean else '', 'quantity':int(qty[0])})
            else:
                m=re.search(r'INBR[A-Z0-9]+',text)
                if m: shipment=m[0]
                m=re.search(r'Total\s+(\d+)',text)
                if m: declared=int(m[1])
                for table in page.extract_tables():
                    for r in table:
                        if len(r)<7: continue
                        skus=str(r[1] or '').strip().splitlines()
                        if not skus or not all(re.fullmatch(r'[A-Za-z]+\d+',sku) for sku in skus): continue
                        quantities=str(r[-1] or '').strip().splitlines()
                        if len(skus)!=len(quantities): raise ValueError('SKU e quantidade não conferem no PDF Shopee.')
                        for sku,qty in zip(skus,quantities): out.append({'sku':sku,'ean':'','quantity':number(qty)})
    if not shipment or not out or declared is None: raise ValueError('PDF Full não reconhecido ou sem total de controle.')
    if sum(r['quantity'] for r in out)!=declared: raise ValueError('A soma das quantidades não confere com o total do PDF. Nenhum item foi aplicado.')
    return {'rows':out,'shipment':shipment,'total':declared,'channel':channel}

if __name__ == '__main__':
    try:
        kind,file=sys.argv[1:3]
        result=full(file,kind) if kind in ('meli','shopee') else spreadsheet(file,kind)
        print(json.dumps(result,ensure_ascii=False,allow_nan=False))
    except Exception as e:
        print(json.dumps({'error':str(e)},ensure_ascii=False));sys.exit(1)
